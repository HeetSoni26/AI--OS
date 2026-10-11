from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

ROOT = Path(r"C:\Users\Heet\OneDrive\Desktop\AI OS")
SOURCE = Path(r"D:\Downloads\AI_Automation_Agency_OS_SRS.docx")
DOCX = ROOT / "AI_OS_Complete_Roadmap_and_12_Build_Prompts.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ID_RE = re.compile(r"\b(?:FR-[A-Z]{2}-\d{3}|NFR-(?:P|S|R|U|PO|M)-\d{3}|SEC-(?:A|Z|E|API|I|C)-\d{3})\b")


def extract_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return "\n".join("".join(node.text or "" for node in p.iter(W + "t")) for p in root.iter(W + "p"))


def extract_source_ids(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    text = "\n".join("".join(node.text or "" for node in p.iter(W + "t")) for p in root.iter(W + "p"))
    return set(ID_RE.findall(text))


def main() -> None:
    required_files = [
        "00_README.md",
        "01_AI_OS_Master_Roadmap.md",
        "02_AI_OS_12_Detailed_Build_Prompts.md",
        "03_SRS_Coverage_and_Traceability_Matrix.md",
        "04_Architecture_Quality_and_Release_Gates.md",
        "AI_OS_Complete_Roadmap_and_12_Build_Prompts.docx",
    ]
    missing = [name for name in required_files if not (ROOT / name).exists()]
    if missing:
        raise SystemExit(f"Missing required deliverables: {missing}")

    roadmap = (ROOT / "01_AI_OS_Master_Roadmap.md").read_text(encoding="utf-8")
    prompts = (ROOT / "02_AI_OS_12_Detailed_Build_Prompts.md").read_text(encoding="utf-8")
    matrix = (ROOT / "03_SRS_Coverage_and_Traceability_Matrix.md").read_text(encoding="utf-8")
    gates = (ROOT / "04_Architecture_Quality_and_Release_Gates.md").read_text(encoding="utf-8")
    readme = (ROOT / "00_README.md").read_text(encoding="utf-8")

    parts = re.findall(r"^## Part (\d+) -", roadmap, flags=re.MULTILINE)
    if parts != [str(i) for i in range(1, 13)]:
        raise SystemExit(f"Roadmap parts are not exactly 1-12 in order: {parts}")
    prompt_numbers = re.findall(r"^## Prompt (\d+) -", prompts, flags=re.MULTILINE)
    if prompt_numbers != [str(i) for i in range(1, 13)]:
        raise SystemExit(f"Prompts are not exactly 1-12 in order: {prompt_numbers}")

    source_ids = extract_source_ids(SOURCE)
    matrix_ids = set(ID_RE.findall(matrix))
    missing_ids = sorted(source_ids - matrix_ids)
    extra_ids = sorted(matrix_ids - source_ids)
    if len(source_ids) != 246 or len(matrix_ids) != 246 or missing_ids or extra_ids:
        raise SystemExit(f"Requirement coverage failed: source={len(source_ids)}, matrix={len(matrix_ids)}, missing={missing_ids}, extra={extra_ids}")

    for part in range(1, 13):
        if f"Part {part}" not in prompts:
            raise SystemExit(f"Prompt handoff missing Part {part}")
    for phrase in ["246/246", "Recommended Addition", "post-MVP", "OpenAPI", "realtime", "webhook", "WCAG 2.1 AA", "RPO", "RTO"]:
        if phrase not in (readme + roadmap + prompts + matrix + gates):
            raise SystemExit(f"Expected package phrase missing: {phrase}")

    document = Document(DOCX)
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name and paragraph.style.name.startswith("Heading")
    ]
    if len(document.sections) < 3:
        raise SystemExit(f"DOCX expected cover/front-matter/body sections, found {len(document.sections)}")
    if len(headings) < 100:
        raise SystemExit(f"DOCX heading count unexpectedly low: {len(headings)}")

    docx_text = extract_docx_text(DOCX)
    for phrase in [
        "Complete Implementation Roadmap",
        "Prompt 1",
        "Prompt 12",
        "246 unique numbered requirements",
        "SRS Coverage and Traceability",
        "Architecture, Quality, and Release Gates",
        "Table of Contents",
    ]:
        if phrase not in docx_text:
            raise SystemExit(f"DOCX content phrase missing: {phrase}")

    with zipfile.ZipFile(DOCX) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
        footer_names = [name for name in archive.namelist() if re.fullmatch(r"word/footer\d+\.xml", name)]
        footer_text = "\n".join(archive.read(name).decode("utf-8") for name in footer_names)
    if "TOC" not in document_xml or 'w:updateFields w:val="true"' not in settings_xml:
        raise SystemExit("DOCX TOC field/updateFields validation failed")
    if "PAGE \\* ROMAN \\* MERGEFORMAT" not in footer_text or "PAGE \\* arabic \\* MERGEFORMAT" not in footer_text:
        raise SystemExit("DOCX footer page-number format validation failed")
    if "<w:pgNumType/>" in document_xml:
        raise SystemExit("Empty cover pgNumType remains")
    if re.search(r'<w:rFonts[^>]+(?:w:ascii|w:hAnsi)="JetBrains Mono"', document_xml):
        raise SystemExit("Non-portable JetBrains Mono run font remains in DOCX")

    print("PASS: required files present")
    print(f"PASS: roadmap parts = {len(parts)} (exactly 1-12)")
    print(f"PASS: detailed prompts = {len(prompt_numbers)} (exactly 1-12)")
    print(f"PASS: numbered SRS coverage = {len(matrix_ids)}/246, missing = 0")
    print(f"PASS: DOCX sections = {len(document.sections)}, headings = {len(headings)}, tables = {len(document.tables)}")
    print("PASS: TOC field, updateFields, Roman/Arabic page fields, cover cleanup, and font portability checks")


if __name__ == "__main__":
    main()
